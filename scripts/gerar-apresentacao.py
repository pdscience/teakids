"""Gera a apresentacao institucional do TEA Kids (PowerPoint).

Uso: python3 scripts/gerar-apresentacao.py
Saida: docs/apresentacao-tea-kids.pptx
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BLUE = RGBColor(0x4D, 0x94, 0xFF)
NAVY = RGBColor(0x0E, 0x1B, 0x33)
INK = RGBColor(0x22, 0x2B, 0x3A)
MUTED = RGBColor(0x5A, 0x66, 0x78)
LIGHT = RGBColor(0xF4, 0xF7, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_BG = RGBColor(0x07, 0x0C, 0x1B)
GREEN = RGBColor(0x07, 0xA6, 0x5A)

WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)

prs = Presentation()
prs.slide_width = WIDTH
prs.slide_height = HEIGHT
BLANK = prs.slide_layouts[6]


def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, HEIGHT)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    return s


def text(s, x, y, w, h, value, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0, font="Segoe UI"):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = spacing
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = font
    return box


def bullets(s, x, y, w, h, items, size=17, color=INK, gap=10, bullet_color=BLUE):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.line_spacing = 1.15
        p.space_after = Pt(gap)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = "Segoe UI"
    return box


def header(s, title, lead=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.62), Inches(0.12), Inches(0.62))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(s, Inches(1.0), Inches(0.55), Inches(11.6), Inches(0.7), title, size=30, bold=True, color=NAVY)
    if lead:
        text(s, Inches(1.0), Inches(1.28), Inches(11.4), Inches(0.5), lead, size=16, color=MUTED)


def footer(s, page):
    text(s, Inches(0.7), Inches(6.95), Inches(8), Inches(0.35), "TEA Kids  |  Comunicação assistida", size=10, color=MUTED)
    text(s, Inches(11.6), Inches(6.95), Inches(1.0), Inches(0.35), str(page), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, icon, title, body, icon_size=26, title_size=17, body_size=12.5):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(0xDC, 0xE4, 0xF2)
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    shape.adjustments[0] = 0.06
    text(s, x + Inches(0.22), y + Inches(0.18), w - Inches(0.44), Inches(0.5), icon, size=icon_size, color=BLUE)
    text(s, x + Inches(0.22), y + Inches(0.78), w - Inches(0.44), Inches(0.4), title, size=title_size, bold=True, color=NAVY)
    text(s, x + Inches(0.22), y + Inches(1.18), w - Inches(0.44), h - Inches(1.35), body, size=body_size, color=MUTED, spacing=1.12)


def callout(s, x, y, w, h, value, size=16, fill=BLUE, color=WHITE, bold=True):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.adjustments[0] = 0.12
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.28)
    tf.margin_right = Inches(0.28)
    p = tf.paragraphs[0]
    p.text = value
    p.line_spacing = 1.15
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Segoe UI"


def stat(s, x, y, w, value, label):
    text(s, x, y, w, Inches(0.7), value, size=34, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.72), w, Inches(0.6), label, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- 1. Capa
s = slide(DARK_BG)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.05), WIDTH, Inches(0.16))
band.fill.solid()
band.fill.fore_color.rgb = BLUE
band.line.fill.background()
band.shadow.inherit = False
text(s, Inches(0.9), Inches(1.85), Inches(11.5), Inches(1.1), "TEA Kids", size=54, bold=True, color=WHITE)
text(s, Inches(0.9), Inches(3.0), Inches(11.0), Inches(0.9),
     "Comunicação assistida para crianças no espectro autista", size=26, color=BLUE)
text(s, Inches(0.9), Inches(3.85), Inches(11.0), Inches(0.6),
     "Um mundo mais azul, um mundo com mais inclusão", size=18, color=RGBColor(0xC5, 0xD2, 0xE6))
text(s, Inches(0.9), Inches(5.45), Inches(11.0), Inches(0.5),
     "Apresentação institucional", size=14, color=RGBColor(0x9A, 0xAA, 0xC2))
text(s, Inches(0.9), Inches(5.85), Inches(11.0), Inches(0.5),
     "Aplicativo web instalável  |  Versão Android  |  Português, Inglês e Espanhol",
     size=14, color=RGBColor(0x9A, 0xAA, 0xC2))

# ---------------------------------------------------------------- 2. O desafio
s = slide()
header(s, "O desafio de concluir uma mensagem",
       "Por que a comunicação da criança trava a rotina inteira")
bullets(s, Inches(1.0), Inches(2.1), Inches(6.6), Inches(3.4), [
    "Comunicar uma necessidade é difícil — e isso trava tudo o que vem depois: comida, sono, dor e transição de atividade.",
    "Frases prontas ajudam, mas não cobrem todas as situações reais do dia a dia.",
    "Rotina sem visualização gera ansiedade, resistência e atrasos.",
    "Família e equipe de acompanhamento precisam enxergar o progresso, não apenas o esforço.",
], size=16.5)
callout(s, Inches(8.1), Inches(2.1), Inches(4.3), Inches(3.0),
        "Sem uma ferramenta que reúna comunicação, rotina e registro, cada dia vira improviso.",
        size=17, fill=NAVY)
footer(s, 2)

# ---------------------------------------------------------------- 3. Proposta
s = slide()
header(s, "A proposta: um app, cinco experiências curtas",
       "Cada área foi desenhada para durar poucos minutos e funcionar com toques simples")
xs = [Inches(0.7), Inches(3.25), Inches(5.8), Inches(8.35), Inches(10.9)]
data = [
    ("💬", "Conversar", "Toca no cartão e a frase é falada em voz alta."),
    ("🧩", "Montar frase", "Combina palavras quando a frase ainda não existe."),
    ("📅", "Rotinas", "Checklist visual de manhã, tarde e noite."),
    ("✏️", "Escrever", "Digita e ouve; explora o alfabeto."),
    ("🎮", "Interação", "Aprende com Quiz: formas, cores, animais e mais."),
]
for x, (icon, title, body) in zip(xs, data):
    card(s, x, Inches(2.15), Inches(2.25), Inches(2.65), icon, title, body, icon_size=28, title_size=15, body_size=11.5)
callout(s, Inches(0.7), Inches(5.25), Inches(11.9), Inches(0.95),
        "A criança usa sozinha: tocar, ouvir, marcar. O adulto acompanha por uma área separada.",
        size=16.5, fill=BLUE)
footer(s, 3)

# ---------------------------------------------------------------- 4. Jornada
s = slide()
header(s, "Como a criança usa — 5 a 10 minutos por sessão",
       "Uma rotina simples, previsível e do tamanho da atenção da criança")
steps = [
    ("1", "Abrir o app", "Vê as estrelas conquistadas no dia."),
    ("2", "Conversar", "Pede, agradece, comenta."),
    ("3", "Marcar a rotina", "Toca no que já fez."),
    ("4", "Brincar", "Quiz de formas, cores e palavras."),
]
for index, (num, title, body) in enumerate(steps):
    x = Inches(0.7 + index * 3.05)
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(2.3), Inches(0.6), Inches(0.6))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = num
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = WHITE
    text(s, x, Inches(3.05), Inches(2.7), Inches(0.4), title, size=16, bold=True, color=NAVY)
    text(s, x, Inches(3.45), Inches(2.6), Inches(0.9), body, size=12.5, color=MUTED)
    if index < len(steps) - 1:
        text(s, x + Inches(2.62), Inches(2.38), Inches(0.4), Inches(0.4), "→", size=20, bold=True, color=BLUE)
callout(s, Inches(0.7), Inches(4.85), Inches(11.9), Inches(1.05),
        "A área dos responsáveis fica atrás de uma confirmação simples para adultos — a criança navega sozinha.",
        size=16.5, fill=NAVY)
footer(s, 4)

# ---------------------------------------------------------------- 5. Conversar
s = slide()
header(s, "Conversar — o momento mais importante do dia",
       "Comunicar uma necessidade em segundos, sem depender de interlocutor disponível")
bullets(s, Inches(1.0), Inches(2.1), Inches(6.4), Inches(3.6), [
    "6 categorias prontas: saudações, necessidades, saúde, atividades, família e locais.",
    "Um toque = a frase é falada em voz alta, no idioma escolhido.",
    "Favoritos e Recentes deixam as frases mais usadas sempre à mão.",
    "Montar frase: quando não existe frase pronta, a criança combina as palavras e ouve.",
], size=16)
card(s, Inches(7.9), Inches(2.1), Inches(4.5), Inches(2.5), "🗣",
     "Benefício para a família",
     "Menos frustração e menos espera: a criança pede o que precisa na hora em que precisa.",
     icon_size=24, title_size=16, body_size=14)
stat(s, Inches(7.9), Inches(4.85), Inches(2.2), "6", "categorias de\ncomunicação")
stat(s, Inches(10.2), Inches(4.85), Inches(2.2), "1 toque", "e a app\nfala")
footer(s, 5)

# ---------------------------------------------------------------- 6. Rotinas
s = slide()
header(s, "Rotinas — previsibilidade que a criança entende",
       "A rotina deixa de ser cobrança e vira um checklist visual que ela acompanha")
bullets(s, Inches(1.0), Inches(2.1), Inches(6.4), Inches(3.6), [
    "Três períodos — Manhã, Tarde e Noite — com 23 atividades já prontas.",
    "Marca com um toque; o progresso fica visível na tela e salvo no aparelho.",
    "Comemoração (som e mensagem) quando o período inteiro é concluído.",
    "O responsável pode acrescentar atividades da casa, sem alterar as padrão.",
], size=16)
card(s, Inches(7.9), Inches(2.1), Inches(4.5), Inches(2.5), "📊",
     "Benefício para o gestor",
     "Cada item concluído é um registro objetivo de adesão à rotina — dia a dia e por período.",
     icon_size=24, title_size=16, body_size=14)
stat(s, Inches(7.9), Inches(4.85), Inches(2.2), "23", "atividades\npadrão")
stat(s, Inches(10.2), Inches(4.85), Inches(2.2), "3", "períodos\npor dia")
footer(s, 6)

# ---------------------------------------------------------------- 7. Escrever + Interação
s = slide()
header(s, "Escrever e Interação — aprender brincando",
       "Alfabetização e vocabulário dentro do mesmo aplicativo, sem material extra")
card(s, Inches(0.7), Inches(2.15), Inches(5.85), Inches(3.3), "✏️", "Escrever",
     "Campo de texto que fala o que foi digitado e um alfabeto de 26 letras, "
     "cada uma com palavra de exemplo e emoji pronunciados ao toque.",
     icon_size=30, title_size=20, body_size=14)
card(s, Inches(6.85), Inches(2.15), Inches(5.85), Inches(3.3), "🎮", "Interação",
     "8 categorias (formas, cores, animais, frutas, alimentos, números, legumes e casa) "
     "com 183 itens e um modo Quiz com feedback imediato.",
     icon_size=30, title_size=20, body_size=14)
callout(s, Inches(0.7), Inches(5.75), Inches(12.0), Inches(0.85),
        "Reforço de vocabulário e alfabetização sem imprimir, recortar nem preparar material.",
        size=16.5, fill=BLUE)
footer(s, 7)

# ---------------------------------------------------------------- 8. Responsáveis
s = slide()
header(s, "Área dos responsáveis — controle do adulto",
       "Tudo o que o adulto precisa, separado da experiência da criança")
items = [
    ("🌐", "Idioma", "Português, Inglês ou Espanhol — troca imediata e salva."),
    ("🔊", "Voz e acessibilidade", "Velocidade, tom, texto maior, alto contraste, menos animações."),
    ("📅", "Histórico (7 dias)", "Uso por período e por módulo, com contadores."),
    ("🃏", "Cartões personalizados", "Frases da família, com emoji e cor."),
    ("📅", "Rotinas personalizadas", "Atividades da casa por período."),
    ("💾", "Dados", "Backup e apagamento seguro com confirmação."),
]
for index, (icon, title, body) in enumerate(items):
    col = index % 3
    row = index // 3
    x = Inches(0.7 + col * 4.05)
    y = Inches(2.15 + row * 1.95)
    card(s, x, y, Inches(3.8), Inches(1.7), icon, title, body, icon_size=22, title_size=15, body_size=11.5)
footer(s, 8)

# ---------------------------------------------------------------- 9. Idioma, voz, acessibilidade
s = slide()
header(s, "Um app que acompanha a criança e a família",
       "Multilíngue de verdade, com voz e acessibilidade já embutidas")
bullets(s, Inches(1.0), Inches(2.1), Inches(6.5), Inches(3.5), [
    "Interface e todo o conteúdo nativo em Português, Inglês e Espanhol.",
    "A voz lê sempre no idioma escolhido; se o aparelho não tiver essa voz, o app usa a mais próxima e avisa — sem travar.",
    "Acessibilidade: texto ampliado, alto contraste, redução de animações e controle da voz.",
    "Funciona sem internet após o primeiro carregamento e pode ser instalado na tela inicial.",
], size=16)
stat(s, Inches(8.0), Inches(2.25), Inches(4.4), "3 idiomas", "mesma experiência\nem cada tela")
card(s, Inches(8.0), Inches(3.65), Inches(4.4), Inches(1.85), "📴",
     "Continuidade",
     "Sem internet no dia a dia: o app continua funcionando no aparelho.",
     icon_size=22, title_size=15, body_size=13)
footer(s, 9)

# ---------------------------------------------------------------- 10. Privacidade
s = slide()
header(s, "Privacidade — os dados ficam com a família",
       "Sem servidor, sem compartilhamento, sem dependência de terceiros")
bullets(s, Inches(1.0), Inches(2.1), Inches(6.5), Inches(3.5), [
    "Progresso, favoritos, ajustes e personalizações são gravados no próprio aparelho.",
    "Nenhuma informação da criança é enviada para servidores ou terceiros.",
    "Backup simples (baixar ou copiar) para guarda e transferência de aparelho.",
    "Apagar todos os dados, a qualquer momento, com confirmação explícita.",
], size=16)
callout(s, Inches(8.0), Inches(2.2), Inches(4.4), Inches(3.2),
        "Em educação e na prática: o app não é mais um cadastro online — é uma ferramenta local, sob controle da família.",
        size=17, fill=NAVY)
footer(s, 10)

# ---------------------------------------------------------------- 11. Ganho do gestor
s = slide()
header(s, "O que o gestor ganha com o TEA Kids",
       "Decisão com base em evidência, não em impressão")
bullets(s, Inches(1.0), Inches(2.05), Inches(6.3), Inches(3.8), [
    "Visão de adesão: rotina concluída por período nos últimos 7 dias.",
    "Visão de engajamento: quais módulos a criança realmente usa.",
    "Conquistas diárias (estrelas) como linguagem simples de evolução.",
    "Menos tempo de preparo: conteúdo, rotina e material já vêm prontos.",
    "Personalização por família sem material impresso.",
], size=15.5)
card(s, Inches(7.6), Inches(2.05), Inches(5.0), Inches(2.35), "📌",
     "Próximos passos sugeridos",
     "Relatórios exportáveis para a escola, metas por criança e acompanhamento compartilhado "
     "entre profissionais.",
     icon_size=22, title_size=16, body_size=13.5)
callout(s, Inches(7.6), Inches(4.6), Inches(5.0), Inches(1.25),
        "Roteiro de implantação sugerido: 1 familiar treina o app; 2 semanas de uso; 1 reunião de leitura dos indicadores.",
        size=14, fill=BLUE)
footer(s, 11)

# ---------------------------------------------------------------- 12. Encerramento
s = slide(DARK_BG)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.75), WIDTH, Inches(0.14))
band.fill.solid()
band.fill.fore_color.rgb = BLUE
band.line.fill.background()
band.shadow.inherit = False
text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.0), "TEA Kids", size=46, bold=True, color=WHITE)
text(s, Inches(0.9), Inches(2.75), Inches(11.0), Inches(0.7),
     "Um mundo mais azul, um mundo com mais inclusão 💙", size=24, color=BLUE)
text(s, Inches(0.9), Inches(3.6), Inches(11.0), Inches(0.6),
     "Próximo passo: demonstração prática com uma família e uma semana de uso acompanhada.",
     size=17, color=RGBColor(0xC5, 0xD2, 0xE6))
text(s, Inches(0.9), Inches(5.15), Inches(11.0), Inches(0.5),
     "Disponível como aplicativo web instalável e em versão Android.",
     size=13, color=RGBColor(0x9A, 0xAA, 0xC2))

prs.save("docs/apresentacao-tea-kids.pptx")
print("ok: docs/apresentacao-tea-kids.pptx")
